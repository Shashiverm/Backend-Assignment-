from openpyxl import load_workbook

def load_excel_data(file_path):
    workbook = load_workbook(filename=file_path)
    sheet = workbook.active
    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        data.append(row)

    return data

def process_customer_data(data):
    customers = []
    for row in data:
        customer = {
            'customer_id': row[0],
            'first_name': row[1],
            'last_name': row[2],
            'phone_number': row[3],
            'monthly_salary': row[4],
            'approved_limit': row[5],
            'current_debt': row[6],
        }
        customers.append(customer)
    return customers

def process_loan_data(data):
    loans = []
    for row in data:
        loan = {
            'loan_id': row[0],
            'customer_id': row[1],
            'loan_amount': row[2],
            'tenure': row[3],
            'interest_rate': row[4],
            'monthly_repayment': row[5],
        }
        loans.append(loan)
    return loans